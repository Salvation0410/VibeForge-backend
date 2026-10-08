from __future__ import annotations

import asyncio
import hashlib
import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, AsyncIterator, Mapping, Protocol, Sequence

from docx import Document
from docx.table import Table
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from ai_service.config import Settings
from ai_service.infrastructure.milvus_knowledge import (
    IndexedChunk,
    IndexedDocument,
    KnowledgeMutationLease,
    KnowledgeStore,
    RebuildFingerprint,
    RebuildPlan,
)
from ai_service.models.embeddings import EmbeddingOutputError, EmbeddingProvider


# Limits also apply after decompression and text extraction, where input byte limits alone do not help.
MAX_EXTRACTED_CHARS = 2_000_000
MAX_CHUNKS = 10_000
MAX_SECTIONS = 10_000
MAX_ZIP_ENTRIES = 1_000
MAX_ZIP_ENTRY_BYTES = 16 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 64 * 1024 * 1024
logger = logging.getLogger(__name__)


class DocumentETLError(RuntimeError):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


@dataclass(frozen=True, slots=True)
class ParsedSection:
    content: str
    source_locator: str


@dataclass(frozen=True, slots=True)
class KnowledgeChunk:
    chunk_id: str
    document_id: str
    document_version: int
    chunk_index: int
    content: str
    source_name: str
    source_locator: str


def _clean(value: str) -> str:
    text = unicodedata.normalize("NFC", value).replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(
        c for c in text
        if c in "\n\t" or (unicodedata.category(c) not in {"Cf", "Cc", "Cs"})
    )
    return "\n".join(re.sub(r"[\t \u00a0]+", " ", line).strip() for line in text.split("\n")).strip()


def _safe_locator(value: str) -> str:
    return re.sub(r"https?://\S+", "[link]", _clean(value), flags=re.IGNORECASE)


def _push_heading(stack: list[tuple[int, str]], level: int, title: str) -> None:
    while stack and stack[-1][0] >= level:
        stack.pop()
    stack.append((level, _safe_locator(title)))


def _append(sections: list[ParsedSection], content: str, locator: str, count: list[int]) -> None:
    cleaned = _clean(content)
    count[0] += len(cleaned)
    if count[0] > MAX_EXTRACTED_CHARS:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
    if cleaned:
        if len(sections) >= MAX_SECTIONS:
            raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
        sections.append(ParsedSection(cleaned, locator))


def _read_text(path: Path) -> str:
    try:
        if path.stat().st_size > MAX_EXTRACTED_CHARS * 4:
            raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
        return path.read_text(encoding="utf-8-sig", errors="strict")
    except UnicodeError:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_INVALID") from None


def _parse_markdown(path: Path) -> list[ParsedSection]:
    sections: list[ParsedSection] = []
    count = [0]
    headings: list[tuple[int, str]] = []
    lines: list[str] = []
    in_code = False

    def flush() -> None:
        if lines:
            _append(sections, "\n".join(lines), " / ".join(title for _, title in headings) or "document", count)
            lines.clear()

    for line_number, line in enumerate(_read_text(path).splitlines(), 1):
        if line_number > MAX_SECTIONS * 10:
            raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
        if re.match(r"^\s*(```|~~~)", line):
            in_code = not in_code
        match = None if in_code else re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if match:
            flush()
            level = len(match.group(1))
            _push_heading(headings, level, match.group(2))
        else:
            lines.append(line)
    flush()
    return sections


def _parse_txt(path: Path) -> list[ParsedSection]:
    lines = _read_text(path).splitlines()
    sections: list[ParsedSection] = []
    count = [0]
    paragraph: list[str] = []
    start = 1
    for number, line in enumerate(lines, 1):
        if line.strip():
            if not paragraph:
                start = number
            paragraph.append(line)
        elif paragraph:
            _append(sections, "\n".join(paragraph), f"lines {start}-{number - 1}", count)
            paragraph.clear()
    if paragraph:
        _append(sections, "\n".join(paragraph), f"lines {start}-{len(lines)}", count)
    return sections


def _check_docx_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES or not any(e.filename == "word/document.xml" for e in entries):
                raise DocumentETLError("KNOWLEDGE_DOCUMENT_INVALID")
            total = 0
            for entry in entries:
                total += entry.file_size
                if (entry.file_size > MAX_ZIP_ENTRY_BYTES or total > MAX_ZIP_TOTAL_BYTES
                        or entry.flag_bits & 1 or "vbaproject.bin" in entry.filename.lower()):
                    raise DocumentETLError("KNOWLEDGE_DOCUMENT_INVALID")
    except DocumentETLError:
        raise
    except Exception:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_INVALID") from None


def _parse_docx(path: Path) -> list[ParsedSection]:
    _check_docx_archive(path)
    sections: list[ParsedSection] = []
    count = [0]
    title = ""
    headings: list[tuple[int, str]] = []

    def read_table(table: Table, locator: str, depth: int) -> None:
        if depth > 8:
            raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
        seen_cells: set[object] = set()
        for row_index, row in enumerate(table.rows, 1):
            if row_index > MAX_SECTIONS:
                raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
            row_locator = f"{locator} / 行 {row_index}"
            texts: list[str] = []

            def flush_texts() -> None:
                if texts:
                    _append(sections, " | ".join(texts), row_locator, count)
                    texts.clear()

            for column_index, cell in enumerate(row.cells, 1):
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                nested_index = 0
                for item in cell.iter_inner_content():
                    if isinstance(item, Table):
                        flush_texts()
                        nested_index += 1
                        read_table(
                            item,
                            f"{row_locator} / 列 {column_index} / 嵌套表格 {nested_index}",
                            depth + 1,
                        )
                    elif item.text.strip():
                        texts.append(item.text)
            flush_texts()

    try:
        doc = Document(path)
        paragraph_index = 0
        table_index = 0
        for block_index, item in enumerate(doc.iter_inner_content(), 1):
            if block_index > MAX_SECTIONS:
                raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
            if isinstance(item, Table):
                table_index += 1
                prefix = [*([title] if title else []), *(text for _, text in headings)]
                read_table(item, " / ".join([*prefix, f"表格 {table_index}"]), 0)
                continue
            paragraph_index += 1
            style = item.style.name or ""
            if style == "Title":
                title = _safe_locator(item.text)
                headings.clear()
            elif match := re.fullmatch(r"Heading ([1-6])", style):
                level = int(match.group(1))
                _push_heading(headings, level, item.text)
            else:
                locator = " / ".join(
                    [*([title] if title else []), *(text for _, text in headings), f"paragraph {paragraph_index}"]
                )
                _append(sections, item.text, locator, count)
    except DocumentETLError:
        raise
    except Exception:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_INVALID") from None
    return sections


def _parse_pdf(path: Path) -> list[ParsedSection]:
    sections: list[ParsedSection] = []
    count = [0]
    try:
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            raise DocumentETLError("KNOWLEDGE_DOCUMENT_ENCRYPTED")
        for page_number, page in enumerate(reader.pages, 1):
            if page_number > MAX_SECTIONS:
                raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
            _append(sections, page.extract_text() or "", f"page {page_number}", count)
    except DocumentETLError:
        raise
    except Exception:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_INVALID") from None
    return sections


def parse_document(path: Path, file_type: str) -> tuple[ParsedSection, ...]:
    parsers = {"md": _parse_markdown, "txt": _parse_txt, "docx": _parse_docx, "pdf": _parse_pdf}
    parser = parsers.get(file_type.lower().lstrip(".")) if isinstance(file_type, str) else None
    if parser is None:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_UNSUPPORTED")
    try:
        sections = parser(Path(path))
    except DocumentETLError:
        raise
    except (OSError, PermissionError):
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_INVALID") from None
    if not sections:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_EMPTY")
    return tuple(sections)


def split_sections(
    sections: tuple[ParsedSection, ...] | list[ParsedSection], *,
    document_id: str, document_version: int, source_name: str, settings: Settings,
) -> tuple[KnowledgeChunk, ...]:
    size, overlap = settings.rag_chunk_size, settings.rag_chunk_overlap
    if not isinstance(size, int) or not isinstance(overlap, int) or size < 1 or overlap < 0 or overlap >= size:
        raise DocumentETLError("KNOWLEDGE_SPLIT_INVALID")
    if not sections or not any(section.content.strip() for section in sections):
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_EMPTY")
    if sum(len(section.content) for section in sections) > MAX_EXTRACTED_CHARS:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
    if (not re.fullmatch(r"[A-Za-z0-9_-]{1,128}", str(document_id))
            or not isinstance(document_version, int) or document_version < 1
            or not isinstance(source_name, str) or not source_name or len(source_name) > 255
            or "://" in source_name or "?" in source_name or "#" in source_name
            or "/" in source_name or "\\" in source_name
            or any(ord(c) < 32 for c in source_name)):
        raise DocumentETLError("KNOWLEDGE_SPLIT_INVALID")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size, chunk_overlap=overlap,
        separators=["\n\n", "\n", "。", "！", "？", "；", " ", ""],
    )
    chunks: list[KnowledgeChunk] = []
    for section in sections:
        if len(chunks) == MAX_CHUNKS and section.content.strip():
            raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
        for content in splitter.split_text(section.content):
            if not content.strip():
                continue
            if len(content) > size:
                raise DocumentETLError("KNOWLEDGE_SPLIT_INVALID")
            if len(chunks) >= MAX_CHUNKS:
                raise DocumentETLError("KNOWLEDGE_DOCUMENT_TOO_LARGE")
            index = len(chunks)
            chunks.append(KnowledgeChunk(
                chunk_id=f"{document_id}:{document_version}:{index}",
                document_id=str(document_id), document_version=document_version,
                chunk_index=index, content=content, source_name=source_name,
                source_locator=_safe_locator(section.source_locator),
            ))
    if not chunks:
        raise DocumentETLError("KNOWLEDGE_DOCUMENT_EMPTY")
    return tuple(chunks)


class _DownloadedFile(Protocol):
    path: Path
    async def __aenter__(self) -> "_DownloadedFile": ...
    async def __aexit__(self, *args: object) -> None: ...


async def parse_and_split_download(
    downloaded: _DownloadedFile, *, file_type: str, document_id: str,
    document_version: int, source_name: str, settings: Settings,
) -> tuple[KnowledgeChunk, ...]:
    def parse_and_split() -> tuple[KnowledgeChunk, ...]:
        return split_sections(
            parse_document(downloaded.path, file_type), document_id=document_id,
            document_version=document_version, source_name=source_name, settings=settings,
        )

    async with downloaded:
        worker = asyncio.create_task(asyncio.to_thread(parse_and_split))
        try:
            return await asyncio.shield(worker)
        except asyncio.CancelledError:
            # The temp file belongs to the worker until parsing and splitting both finish.
            while not worker.done():
                try:
                    await asyncio.shield(worker)
                except asyncio.CancelledError:
                    continue
                except Exception:
                    break
            try:
                worker.result()
            except Exception:
                pass
            raise


class KnowledgeEtlService:
    """Compose download, parsing, embedding, and the version-aware knowledge store."""

    def __init__(
        self,
        settings: Settings,
        downloader: object,
        embeddings: EmbeddingProvider,
        store: KnowledgeStore,
        semaphore: asyncio.Semaphore | None = None,
    ) -> None:
        self._settings = settings
        self._downloader = downloader
        self._embeddings = embeddings
        self._store = store
        self._semaphore = semaphore or asyncio.Semaphore(
            settings.rag_etl_max_concurrency
        )

    async def index(
        self, *, document_id: str, document_version: int, file_name: str,
        file_type: str, signed_url: str, sha256: str, etl_version: str,
        lease: KnowledgeMutationLease,
    ):
        async with self._semaphore:
            return await self._index(
                document_id=document_id,
                document_version=document_version,
                file_name=file_name,
                file_type=file_type,
                signed_url=signed_url,
                sha256=sha256,
                etl_version=etl_version,
                lease=lease,
            )

    async def _index(
        self, *, document_id: str, document_version: int, file_name: str,
        file_type: str, signed_url: str, sha256: str, etl_version: str,
        lease: KnowledgeMutationLease,
    ):
        logger.info("Knowledge INDEX downloading: documentId=%s", document_id)
        downloaded = await self._downloader.download(
            signed_url,
            expected_sha256=sha256,
            max_bytes=self._settings.rag_download_max_bytes,
        )
        logger.info("Knowledge INDEX parsing: documentId=%s fileType=%s", document_id, file_type)
        chunks = await parse_and_split_download(
            downloaded,
            file_type=file_type,
            document_id=document_id,
            document_version=document_version,
            source_name=file_name,
            settings=self._settings,
        )
        dimension = self._settings.rag_embedding_dimension
        if len(chunks) * dimension > self._settings.rag_max_embedding_elements:
            raise EmbeddingOutputError("KNOWLEDGE_EMBEDDING_BUDGET_EXCEEDED")
        logger.info("Knowledge INDEX embedding: documentId=%s chunks=%s", document_id, len(chunks))
        vectors = await self._embeddings.embed_documents(
            [chunk.content for chunk in chunks],
            max_elements=self._settings.rag_max_embedding_elements,
        )
        if len(vectors) != len(chunks):
            raise EmbeddingOutputError("KNOWLEDGE_EMBEDDING_COUNT_MISMATCH")
        try:
            wrong_dimension = any(len(vector) != dimension for vector in vectors)
        except TypeError:
            raise EmbeddingOutputError("KNOWLEDGE_EMBEDDING_INVALID_OUTPUT") from None
        if wrong_dimension:
            raise EmbeddingOutputError("KNOWLEDGE_EMBEDDING_DIMENSION_MISMATCH")
        indexed_chunks = tuple(
            IndexedChunk(
                chunk_id=chunk.chunk_id,
                chunk_index=chunk.chunk_index,
                content=chunk.content,
                source_locator=chunk.source_locator,
                content_hash=hashlib.sha256(chunk.content.encode("utf-8")).hexdigest(),
                embedding=vector,
            )
            for chunk, vector in zip(chunks, vectors, strict=True)
        )
        document = IndexedDocument(
            document_id=document_id,
            document_version=document_version,
            etl_version=etl_version,
            embedding_model_version=self._settings.rag_embedding_model,
            file_name=file_name,
            file_type=file_type.lower(),
            content_hash=sha256,
            chunks=indexed_chunks,
        )
        logger.info("Knowledge INDEX storing: documentId=%s version=%s", document_id, document_version)
        return await self._store.upsert_document_version(document, lease=lease)

    async def delete(
        self, *, document_id: str, document_version: int,
        lease: KnowledgeMutationLease,
    ) -> None:
        await self._store.delete_document(
            document_id, document_version, lease=lease
        )

    async def rebuild(
        self, *, documents: Sequence[Any], etl_version: str,
        lease: KnowledgeMutationLease,
    ):
        async with self._semaphore:
            total_chunks = 0
            total_text_bytes = 0

            def field(document: Any, name: str) -> Any:
                if isinstance(document, Mapping):
                    camel = name.split("_")[0] + "".join(
                        part.capitalize() for part in name.split("_")[1:]
                    )
                    return document[camel]
                return getattr(document, name)

            fingerprint_builder = RebuildFingerprint(
                self._settings.milvus_collection_alias, etl_version
            )
            for item in documents:
                fingerprint_builder.add_document(
                    document_id=field(item, "document_id"),
                    document_version=field(item, "document_version"),
                    file_name=field(item, "file_name"),
                    file_type=field(item, "file_type"),
                    content_hash=field(item, "sha256"),
                )
            fingerprint = fingerprint_builder.hexdigest()
            plan = RebuildPlan(
                embedding_model_version=self._settings.rag_embedding_model,
                embedding_dimension=self._settings.rag_embedding_dimension,
                etl_version=etl_version,
                document_count=len(documents),
                document_fingerprint=fingerprint,
            )

            async def indexed_documents() -> AsyncIterator[IndexedDocument]:
                nonlocal total_chunks, total_text_bytes
                prepared: list[tuple[Any, tuple[KnowledgeChunk, ...]]] = []
                for item in documents:
                    document_id = field(item, "document_id")
                    document_version = field(item, "document_version")
                    file_name = field(item, "file_name")
                    file_type = field(item, "file_type")
                    signed_url = field(item, "signed_url")
                    sha256 = field(item, "sha256")
                    downloaded = await self._downloader.download(
                        signed_url,
                        expected_sha256=sha256,
                        max_bytes=self._settings.rag_download_max_bytes,
                    )
                    chunks = await parse_and_split_download(
                        downloaded,
                        file_type=file_type,
                        document_id=document_id,
                        document_version=document_version,
                        source_name=file_name,
                        settings=self._settings,
                    )
                    total_chunks += len(chunks)
                    if total_chunks > self._settings.rag_rebuild_max_chunks:
                        raise DocumentETLError("KNOWLEDGE_REBUILD_TOO_MANY_CHUNKS")
                    if (
                        total_chunks * self._settings.rag_embedding_dimension
                        > self._settings.rag_max_embedding_elements
                    ):
                        raise EmbeddingOutputError(
                            "KNOWLEDGE_EMBEDDING_BUDGET_EXCEEDED"
                        )
                    total_text_bytes += sum(
                        len(chunk.content.encode("utf-8")) for chunk in chunks
                    )
                    if total_text_bytes > self._settings.rag_rebuild_max_text_bytes:
                        raise DocumentETLError(
                            "KNOWLEDGE_REBUILD_TEXT_BUDGET_EXCEEDED"
                        )
                    prepared.append((item, chunks))

                embedded_chunks = 0
                for item, chunks in prepared:
                    document_id = field(item, "document_id")
                    document_version = field(item, "document_version")
                    file_name = field(item, "file_name")
                    file_type = field(item, "file_type")
                    sha256 = field(item, "sha256")
                    remaining = self._settings.rag_max_embedding_elements - (
                        embedded_chunks * self._settings.rag_embedding_dimension
                    )
                    vectors = await self._embeddings.embed_documents(
                        [chunk.content for chunk in chunks],
                        max_elements=remaining,
                    )
                    if len(vectors) != len(chunks):
                        raise EmbeddingOutputError(
                            "KNOWLEDGE_EMBEDDING_COUNT_MISMATCH"
                        )
                    try:
                        wrong_dimension = any(
                            len(vector) != self._settings.rag_embedding_dimension
                            for vector in vectors
                        )
                    except TypeError:
                        raise EmbeddingOutputError(
                            "KNOWLEDGE_EMBEDDING_INVALID_OUTPUT"
                        ) from None
                    if wrong_dimension:
                        raise EmbeddingOutputError(
                            "KNOWLEDGE_EMBEDDING_DIMENSION_MISMATCH"
                        )
                    embedded_chunks += len(chunks)
                    try:
                        indexed_chunks = tuple(
                            IndexedChunk(
                                chunk_id=chunk.chunk_id,
                                chunk_index=chunk.chunk_index,
                                content=chunk.content,
                                source_locator=chunk.source_locator,
                                content_hash=hashlib.sha256(
                                    chunk.content.encode("utf-8")
                                ).hexdigest(),
                                embedding=vector,
                            )
                            for chunk, vector in zip(chunks, vectors, strict=True)
                        )
                    except ValueError:
                        raise EmbeddingOutputError(
                            "KNOWLEDGE_EMBEDDING_INVALID_OUTPUT"
                        ) from None
                    yield IndexedDocument(
                        document_id=document_id,
                        document_version=document_version,
                        etl_version=etl_version,
                        embedding_model_version=self._settings.rag_embedding_model,
                        file_name=file_name,
                        file_type=file_type.lower(),
                        content_hash=sha256,
                        chunks=indexed_chunks,
                    )

            return await self._store.rebuild_collection(
                indexed_documents(), lease=lease, plan=plan
            )

    async def ping(self) -> bool:
        return await self._store.ping()
