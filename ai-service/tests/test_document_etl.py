from __future__ import annotations

from dataclasses import FrozenInstanceError
from pathlib import Path
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from ai_service.orchestration.document_etl import (
    DocumentETLError,
    ParsedSection,
    parse_document,
    split_sections,
    parse_and_split_download,
)


FIXTURES = Path(__file__).parent / "fixtures" / "knowledge"


def test_markdown_heading_paths_and_immutable_sections():
    sections = parse_document(FIXTURES / "sample.md", "md")
    assert any(s.source_locator == "产品指南 / 退款" and "七天" in s.content for s in sections)
    assert any(s.source_locator == "产品指南 / 退款 / 时限" for s in sections)
    with pytest.raises(FrozenInstanceError):
        sections[0].content = "changed"


def test_txt_line_ranges_and_unicode_cleanup(tmp_path):
    sections = parse_document(FIXTURES / "sample.txt", "txt")
    assert sections[0].source_locator.startswith("lines 1-")
    assert sections[-1].source_locator == "lines 4-4"
    assert "第四行" in " ".join(s.content for s in sections)
    p = tmp_path / "unicode.txt"
    p.write_text("ＡＢＣ\u200b  内容\r\n第二行", encoding="utf-8")
    cleaned = parse_document(p, "txt")
    assert "ＡＢＣ 内容" in cleaned[0].content
    assert "\u200b" not in cleaned[0].content


@pytest.mark.parametrize("file_type", ["md", "txt"])
def test_empty_text_document_rejected(tmp_path, file_type):
    p = tmp_path / f"empty.{file_type}"
    p.write_text(" \n", encoding="utf-8")
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_EMPTY"):
        parse_document(p, file_type)


def test_signed_link_in_heading_is_redacted_from_locator(tmp_path):
    p = tmp_path / "link.md"
    p.write_text("# https://oss.example.test/doc?Signature=secret\n\nContent", encoding="utf-8")
    sections = parse_document(p, "md")
    assert sections[0].source_locator == "[link]"


def test_docx_title_heading_and_paragraph_locator(tmp_path):
    p = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("售后政策", style="Title")
    doc.add_heading("退款", level=1)
    doc.add_paragraph("订单七天内可退款。")
    doc.save(p)
    sections = parse_document(p, "docx")
    assert any(s.source_locator == "售后政策 / 退款 / paragraph 3" for s in sections)


def test_pdf_page_locator_and_encryption(tmp_path):
    p = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with p.open("wb") as out:
        writer.write(out)
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_EMPTY"):
        parse_document(p, "pdf")
    writer.encrypt("secret")
    with p.open("wb") as out:
        writer.write(out)
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_ENCRYPTED"):
        parse_document(p, "pdf")


def test_pdf_extracted_text_keeps_page_number(tmp_path):
    p = tmp_path / "text.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Refund policy) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with p.open("wb") as out:
        writer.write(out)
    sections = parse_document(p, "pdf")
    assert sections[0].source_locator == "page 1"
    assert "Refund policy" in sections[0].content


def test_corrupt_unsupported_and_zip_bomb_rejected(tmp_path):
    p = tmp_path / "bad.pdf"
    p.write_bytes(b"not a pdf")
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_INVALID") as caught:
        parse_document(p, "pdf")
    assert str(p) not in str(caught.value)
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_UNSUPPORTED"):
        parse_document(p, "exe")
    docx = tmp_path / "bad.docx"
    docx.write_bytes(b"not a docx")
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_INVALID"):
        parse_document(docx, "docx")


def test_docx_macro_and_expanded_zip_limit_rejected(tmp_path):
    p = tmp_path / "macro.docx"
    Document().save(p)
    with zipfile.ZipFile(p, "a") as archive:
        archive.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_INVALID"):
        parse_document(p, "docx")
    p = tmp_path / "expanded.docx"
    Document().save(p)
    with zipfile.ZipFile(p, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("extra.bin", b"0" * (16 * 1024 * 1024 + 1))
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_INVALID"):
        parse_document(p, "docx")


def test_splitter_chunk_ids_overlap_and_metadata(settings, tmp_path):
    settings.rag_chunk_size = 1000
    settings.rag_chunk_overlap = 150
    p = tmp_path / "long.txt"
    p.write_text("知识" * 600, encoding="utf-8")
    sections = parse_document(p, "txt")
    chunks = split_sections(
        sections, document_id="42", document_version=3,
        source_name="manual.txt", settings=settings,
    )
    assert len(chunks) > 1
    assert [c.chunk_id for c in chunks] == [f"42:3:{i}" for i in range(len(chunks))]
    assert all(0 < len(c.content) <= 1000 for c in chunks)
    assert chunks[0].content[-150:] == chunks[1].content[:150]
    assert all(c.source_name == "manual.txt" and c.source_locator for c in chunks)
    with pytest.raises(FrozenInstanceError):
        chunks[0].content = "changed"


def test_invalid_overlap_and_empty_document_fail(settings):
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_DOCUMENT_EMPTY"):
        split_sections((), document_id="42", document_version=1, source_name="x", settings=settings)
    settings.rag_chunk_overlap = settings.rag_chunk_size
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_SPLIT_INVALID"):
        split_sections(parse_document(FIXTURES / "sample.md", "md"), document_id="42", document_version=1, source_name="x", settings=settings)


def test_splitter_redacts_caller_supplied_locator_and_rejects_signed_source_name(settings):
    sections = (ParsedSection("Content", "https://oss.example.test/file?Signature=secret"),)
    chunks = split_sections(
        sections, document_id="42", document_version=1,
        source_name="manual.txt", settings=settings,
    )
    assert chunks[0].source_locator == "[link]"
    with pytest.raises(DocumentETLError, match="KNOWLEDGE_SPLIT_INVALID"):
        split_sections(
            sections, document_id="42", document_version=1,
            source_name="https://oss.example.test/file?Signature=secret", settings=settings,
        )


@pytest.mark.asyncio
async def test_pipeline_cleans_download_on_success_and_parse_failure(settings, tmp_path):
    class Downloaded:
        def __init__(self, path):
            self.path = path
        async def __aenter__(self):
            return self
        async def __aexit__(self, *_):
            self.path.unlink(missing_ok=True)

    good = tmp_path / "good.txt"
    good.write_text("客服内容", encoding="utf-8")
    chunks = await parse_and_split_download(
        Downloaded(good), file_type="txt", document_id="42",
        document_version=1, source_name="help.txt", settings=settings,
    )
    assert chunks and not good.exists()
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"bad")
    with pytest.raises(DocumentETLError):
        await parse_and_split_download(
            Downloaded(bad), file_type="pdf", document_id="42",
            document_version=1, source_name="help.pdf", settings=settings,
        )
    assert not bad.exists()
